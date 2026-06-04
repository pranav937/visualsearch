import csv
import os
import subprocess
import sys
import time

# Auto-install and import required libraries
def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    from PIL import Image as PILImage
except ImportError:
    install("Pillow")
    from PIL import Image as PILImage

try:
    import openpyxl
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.utils import get_column_letter
except ImportError:
    install("openpyxl")
    import openpyxl
    from openpyxl.drawing.image import Image as OpenpyxlImage
    from openpyxl.utils import get_column_letter

INPUT_CSV = 'all_data_final_with_images.csv'
OUTPUT_EXCEL = 'all_data_final_with_images.xlsx'
IMAGE_SIZE = 120  # Set identical width and height for all images
CELL_HEIGHT = 100
CELL_WIDTH = 20

def main():
    if not os.path.exists(INPUT_CSV):
        print(f"Error: {INPUT_CSV} not found.")
        return

    print("Initializing Excel workbook...")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Products with Images"

    with open(INPUT_CSV, mode='r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        fieldnames = list(reader.fieldnames)
        
        # We will replace 'Local Image Path' column with the actual image
        # Let's keep the column name, or change it to 'Product Image'
        if 'Local Image Path' in fieldnames:
            img_col_idx = fieldnames.index('Local Image Path')
            fieldnames[img_col_idx] = 'Product Image'
        else:
            fieldnames.append('Product Image')
            img_col_idx = len(fieldnames) - 1
            
        img_col_letter = get_column_letter(img_col_idx + 1)
        
        # Write Headers
        ws.append(fieldnames)
        
        # Set column width
        ws.column_dimensions[img_col_letter].width = CELL_WIDTH

        print(f"Processing data and embedding images (resizing all to {IMAGE_SIZE}x{IMAGE_SIZE})...")
        row_num = 2 # Row 1 is header
        count = 0
        
        for row in reader:
            local_path = row.get('Local Image Path', '').strip()
            
            # Prepare row data, leave image cell blank text
            row_data = []
            for k in reader.fieldnames:
                if k == 'Local Image Path':
                    row_data.append('')  # We will insert the image here
                else:
                    row_data.append(row.get(k, ''))
                    
            ws.append(row_data)
            
            # Set row height
            ws.row_dimensions[row_num].height = CELL_HEIGHT
            
            if local_path and os.path.exists(local_path):
                try:
                    # Insert image into Excel
                    img = OpenpyxlImage(local_path)
                    img.width, img.height = IMAGE_SIZE, IMAGE_SIZE
                    
                    # Position image in the exact cell
                    img.anchor = f"{img_col_letter}{row_num}"
                    ws.add_image(img)
                    
                    if count % 100 == 0:
                        print(f"[{count}] Embedded image from: {local_path}")
                except Exception as e:
                    print(f"Error embedding image {local_path}: {e}")
            
            row_num += 1
            count += 1
            
    # Save Excel file
    print(f"\nSaving Excel file to {OUTPUT_EXCEL}...")
    wb.save(OUTPUT_EXCEL)
    print(f"Finished processing {count} products successfully.")

if __name__ == "__main__":
    main()
